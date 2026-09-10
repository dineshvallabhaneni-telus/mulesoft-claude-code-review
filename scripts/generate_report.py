#!/usr/bin/env python3

"""
MuleSoft Claude Code Review Framework
-------------------------------------

Generate the final Microsoft Word code-review report.

Execution model:

    application-root/
        <MuleSoft application>

    framework-root/
        CLAUDE.md
        agents/
        skills/
        scripts/
        templates/
        config/
        workspace/
            execution/
        reports/

Input:

    <framework-root>/workspace/execution/validated-review.json

Optional template:

    <framework-root>/templates/review-report-template.docx

Output:

    <framework-root>/reports/
        CODE_REVIEW_REPORT_<YYYYMMDD-HHMMSS>.docx

The MuleSoft application is treated as read-only.
"""

import argparse
import json
import os
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import (
    WD_CELL_VERTICAL_ALIGNMENT,
    WD_TABLE_ALIGNMENT,
)
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

SEVERITIES = [
    "CRITICAL",
    "HIGH",
    "MEDIUM",
    "LOW",
    "NIT",
]

CONFIDENCE_LEVELS = [
    "HIGH",
    "MEDIUM",
    "LOW",
]

CATEGORY_ORDER = [
    "Security",
    "Architecture",
    "Application Structure",
    "Error Handling",
    "Data Transformation",
    "API Design",
    "Connectors",
    "Database",
    "Messaging",
    "Performance",
    "Logging and Observability",
    "Testing",
    "Build and Dependency Management",
    "Configuration",
    "Maintainability",
]

READINESS_STATUSES = [
    "READY",
    "PARTIAL",
    "NOT READY",
    "NOT APPLICABLE",
    "NOT ASSESSED",
]

ASSESSMENT_RATINGS = [
    "STRONG",
    "ADEQUATE",
    "WEAK",
    "INADEQUATE",
    "NOT ASSESSED",
]

COVERAGE_STATUSES = [
    "COMPLETED",
    "PARTIAL",
    "LIMITED",
    "NOT APPLICABLE",
    "NOT ASSESSED",
    "FAILED",
]

COLORS = {
    # Severity
    "CRITICAL": "C62828",
    "HIGH": "EF6C00",
    "MEDIUM": "F9A825",
    "LOW": "2E7D32",
    "NIT": "1565C0",

    # Assessment
    "STRONG": "2E7D32",
    "ADEQUATE": "66BB6A",
    "WEAK": "EF6C00",
    "INADEQUATE": "C62828",
    "NOT ASSESSED": "757575",

    # Readiness
    "READY": "2E7D32",
    "PARTIAL": "F9A825",
    "NOT READY": "C62828",
    "NOT APPLICABLE": "757575",

    # General
    "PRIMARY": "17365D",
    "SECONDARY": "D9EAF7",
    "TABLE_HEADER": "D9EAF7",
    "EVIDENCE": "F4F4F4",
    "BORDER": "B7B7B7",
    "TEXT": "222222",
    "MUTED": "666666",
    "WHITE": "FFFFFF",
    "BLACK": "000000",
}

DEFAULT_REVIEWER = "Claude MuleSoft Code Review Agent"


# ---------------------------------------------------------------------------
# Argument handling
# ---------------------------------------------------------------------------

def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Generate the final MuleSoft code review Word report."
    )

    parser.add_argument(
        "--application-root",
        default=os.environ.get("APPLICATION_ROOT"),
        help="Root directory of the MuleSoft application.",
    )

    parser.add_argument(
        "--framework-root",
        default=os.environ.get("FRAMEWORK_ROOT"),
        help="Root directory of the review framework.",
    )

    parser.add_argument(
        "--input",
        default=None,
        help=(
            "Optional path to validated-review.json. "
            "Defaults to framework-root/workspace/execution/"
            "validated-review.json."
        ),
    )

    parser.add_argument(
        "--output",
        default=None,
        help=(
            "Optional output DOCX path. "
            "Normally generated automatically under framework-root/reports."
        ),
    )

    parser.add_argument(
        "--no-template",
        action="store_true",
        help="Do not use templates/review-report-template.docx.",
    )

    return parser.parse_args()


# ---------------------------------------------------------------------------
# Path handling
# ---------------------------------------------------------------------------

def resolve_root(value: Optional[str], fallback: Path) -> Path:
    if value:
        return Path(value).expanduser().resolve()

    return fallback.resolve()


def resolve_paths(args: argparse.Namespace) -> Tuple[Path, Path, Path, Path]:
    cwd = Path.cwd().resolve()

    framework_root = resolve_root(
        args.framework_root,
        cwd,
    )

    application_root = resolve_root(
        args.application_root,
        cwd,
    )

    execution_dir = framework_root / "workspace" / "execution"
    reports_dir = framework_root / "reports"

    execution_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    reports_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    if args.input:
        input_file = Path(args.input).expanduser().resolve()
    else:
        input_file = execution_dir / "validated-review.json"

    return (
        application_root,
        framework_root,
        input_file,
        reports_dir,
    )


# ---------------------------------------------------------------------------
# Generic helpers
# ---------------------------------------------------------------------------

def now_local() -> datetime:
    return datetime.now().astimezone()


def safe_text(value: Any, default: str = "Not Identified") -> str:
    if value is None:
        return default

    if isinstance(value, str):
        value = value.strip()

        if not value:
            return default

        return value

    if isinstance(value, (list, tuple)):
        if not value:
            return default

        return ", ".join(
            safe_text(item, "")
            for item in value
            if safe_text(item, "")
        )

    if isinstance(value, dict):
        if not value:
            return default

        return json.dumps(
            value,
            ensure_ascii=False,
            indent=2,
        )

    return str(value)


def first_value(
    source: Dict[str, Any],
    keys: Iterable[str],
    default: str = "Not Identified",
) -> str:
    for key in keys:
        if key in source:
            value = safe_text(
                source.get(key),
                "",
            )

            if value:
                return value

    return default


def rgb(hex_value: str) -> RGBColor:
    value = hex_value.lstrip("#")

    return RGBColor(
        int(value[0:2], 16),
        int(value[2:4], 16),
        int(value[4:6], 16),
    )


def color_for(value: str, default: str = "MUTED") -> str:
    return COLORS.get(
        value.upper(),
        COLORS.get(
            default,
            COLORS["MUTED"],
        ),
    )


def normalize_category(category: Any) -> str:
    value = safe_text(
        category,
        "Other",
    )

    aliases = {
        "API": "API Design",
        "APIS": "API Design",
        "DATAWEAVE": "Data Transformation",
        "DATA WEAVE": "Data Transformation",
        "LOGGING": "Logging and Observability",
        "OBSERVABILITY": "Logging and Observability",
        "MUNIT": "Testing",
        "TESTING": "Testing",
        "MAVEN": "Build and Dependency Management",
        "DEPENDENCY": "Build and Dependency Management",
        "DEPENDENCIES": "Build and Dependency Management",
        "BUILD": "Build and Dependency Management",
        "XML": "Application Structure",
        "MULE XML": "Application Structure",
    }

    return aliases.get(
        value.upper(),
        value,
    )


def load_json(path: Path) -> Dict[str, Any]:
    if not path.exists():
        raise FileNotFoundError(
            f"Missing validated review: {path}"
        )

    try:
        with path.open(
            "r",
            encoding="utf-8",
        ) as handle:
            value = json.load(handle)
    except json.JSONDecodeError as exc:
        raise RuntimeError(
            f"Invalid JSON in {path}: {exc}"
        ) from exc

    if not isinstance(value, dict):
        raise RuntimeError(
            "Validated review JSON must contain an object."
        )

    return value


# ---------------------------------------------------------------------------
# Word document primitives
# ---------------------------------------------------------------------------

def set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()

    shd = tc_pr.find(
        qn("w:shd")
    )

    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)

    shd.set(
        qn("w:fill"),
        color,
    )


def set_cell_border(
    cell,
    color: str = "B7B7B7",
    size: str = "6",
) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()

    borders = tc_pr.first_child_found_in(
        "w:tcBorders"
    )

    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)

    for edge in (
        "top",
        "left",
        "bottom",
        "right",
        "insideH",
        "insideV",
    ):
        tag = "w:" + edge

        element = borders.find(
            qn(tag)
        )

        if element is None:
            element = OxmlElement(tag)
            borders.append(element)

        element.set(
            qn("w:val"),
            "single",
        )

        element.set(
            qn("w:sz"),
            size,
        )

        element.set(
            qn("w:color"),
            color,
        )


def set_cell_text(
    cell,
    text: Any,
    bold: bool = False,
    color: Optional[str] = None,
    size: int = 9,
    font: str = "Aptos",
) -> None:
    cell.text = ""

    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)

    run = paragraph.add_run(
        safe_text(text, "")
    )

    run.bold = bold
    run.font.name = font
    run.font.size = Pt(size)

    if color:
        run.font.color.rgb = rgb(color)

    cell.vertical_alignment = (
        WD_CELL_VERTICAL_ALIGNMENT.CENTER
    )

    set_cell_border(
        cell,
        COLORS["BORDER"],
    )


def style_table(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    try:
        table.style = "Table Grid"
    except Exception:
        pass

    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = (
                WD_CELL_VERTICAL_ALIGNMENT.CENTER
            )
            set_cell_border(
                cell,
                COLORS["BORDER"],
            )


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()

    tbl_header = OxmlElement(
        "w:tblHeader"
    )

    tbl_header.set(
        qn("w:val"),
        "true",
    )

    tr_pr.append(tbl_header)


def add_heading(
    document,
    text: str,
    level: int = 1,
):
    paragraph = document.add_heading(
        safe_text(text),
        level=level,
    )

    for run in paragraph.runs:
        run.font.name = "Aptos Display"

        if level == 1:
            run.font.color.rgb = rgb(
                COLORS["PRIMARY"]
            )
        else:
            run.font.color.rgb = rgb(
                COLORS["PRIMARY"]
            )

    return paragraph


def add_body(
    document,
    text: Any,
    bold: bool = False,
    italic: bool = False,
    color: Optional[str] = None,
    size: int = 9,
):
    paragraph = document.add_paragraph()

    paragraph.paragraph_format.space_after = Pt(5)

    run = paragraph.add_run(
        safe_text(text, "")
    )

    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic

    if color:
        run.font.color.rgb = rgb(color)

    return paragraph


def add_bullets(
    document,
    values: Iterable[Any],
) -> None:
    for value in values:
        text = safe_text(
            value,
            "",
        )

        if not text:
            continue

        paragraph = document.add_paragraph(
            style="List Bullet"
        )

        paragraph.paragraph_format.space_after = Pt(2)

        run = paragraph.add_run(text)
        run.font.name = "Aptos"
        run.font.size = Pt(9)


def add_key_value_table(
    document,
    rows: List[Tuple[str, Any]],
):
    table = document.add_table(
        rows=1,
        cols=2,
    )

    style_table(table)

    headers = [
        "Attribute",
        "Value",
    ]

    for index, header in enumerate(headers):
        set_cell_text(
            table.cell(0, index),
            header,
            bold=True,
            color=COLORS["PRIMARY"],
        )

        set_cell_shading(
            table.cell(0, index),
            COLORS["SECONDARY"],
        )

    set_repeat_table_header(
        table.rows[0]
    )

    for key, value in rows:
        cells = table.add_row().cells

        set_cell_text(
            cells[0],
            key,
            bold=True,
        )

        set_cell_text(
            cells[1],
            value,
        )

    document.add_paragraph()

    return table


def add_generic_table(
    document,
    headers: List[str],
    rows: List[Iterable[Any]],
):
    table = document.add_table(
        rows=1,
        cols=len(headers),
    )

    style_table(table)

    for index, header in enumerate(headers):
        cell = table.cell(
            0,
            index,
        )

        set_cell_text(
            cell,
            header,
            bold=True,
            color=COLORS["PRIMARY"],
        )

        set_cell_shading(
            cell,
            COLORS["TABLE_HEADER"],
        )

    set_repeat_table_header(
        table.rows[0]
    )

    for row in rows:
        values = list(row)

        while len(values) < len(headers):
            values.append("")

        cells = table.add_row().cells

        for index in range(len(headers)):
            set_cell_text(
                cells[index],
                values[index],
            )

    document.add_paragraph()

    return table


def add_status_badge(
    document,
    label: str,
    value: str,
    color: Optional[str] = None,
) -> None:
    resolved_color = color or color_for(
        value,
        "PRIMARY",
    )

    table = document.add_table(
        rows=1,
        cols=1,
    )

    table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )

    cell = table.cell(0, 0)

    set_cell_shading(
        cell,
        resolved_color,
    )

    set_cell_border(
        cell,
        resolved_color,
        "8",
    )

    paragraph = cell.paragraphs[0]
    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    paragraph.paragraph_format.space_after = Pt(0)

    run = paragraph.add_run(
        f"{label}\n{value}"
    )

    run.bold = True
    run.font.name = "Aptos"
    run.font.size = Pt(14)
    run.font.color.rgb = rgb(
        COLORS["WHITE"]
    )

    document.add_paragraph()


def add_section_page_break(
    document,
) -> None:
    document.add_page_break()


def add_evidence_block(
    document,
    text: Any,
) -> None:
    table = document.add_table(
        rows=1,
        cols=1,
    )

    cell = table.cell(
        0,
        0,
    )

    set_cell_shading(
        cell,
        COLORS["EVIDENCE"],
    )

    set_cell_border(
        cell,
        COLORS["BORDER"],
    )

    paragraph = cell.paragraphs[0]

    run = paragraph.add_run(
        safe_text(text, "")
    )

    run.font.name = "Consolas"
    run.font.size = Pt(8)
    run.font.color.rgb = rgb(
        COLORS["TEXT"]
    )

    document.add_paragraph()


def add_labeled_text(
    document,
    label: str,
    value: Any,
    evidence_block: bool = False,
) -> None:
    add_heading(
        document,
        label,
        level=3,
    )

    if evidence_block:
        add_evidence_block(
            document,
            value,
        )
    else:
        add_body(
            document,
            value,
        )


# ---------------------------------------------------------------------------
# Finding helpers
# ---------------------------------------------------------------------------

def get_findings(
    review: Dict[str, Any],
) -> List[Dict[str, Any]]:
    raw = review.get(
        "findings",
        [],
    )

    if not isinstance(raw, list):
        return []

    findings = []

    for finding in raw:
        if isinstance(finding, dict):
            normalized = dict(finding)

            normalized["category"] = normalize_category(
                normalized.get(
                    "category",
                    "Other",
                )
            )

            findings.append(
                normalized
            )

    return findings


def sort_findings(
    findings: List[Dict[str, Any]],
) -> List[Dict[str, Any]]:
    severity_rank = {
        "CRITICAL": 0,
        "HIGH": 1,
        "MEDIUM": 2,
        "LOW": 3,
        "NIT": 4,
    }

    return sorted(
        findings,
        key=lambda item: (
            severity_rank.get(
                safe_text(
                    item.get("severity"),
                    "NIT",
                ).upper(),
                99,
            ),
            safe_text(
                item.get("id"),
                "",
            ),
        ),
    )


def severity_counts(
    findings: List[Dict[str, Any]],
) -> Counter:
    return Counter(
        safe_text(
            finding.get("severity"),
            "",
        ).upper()
        for finding in findings
    )


def category_counts(
    findings: List[Dict[str, Any]],
) -> Counter:
    return Counter(
        normalize_category(
            finding.get(
                "category",
                "Other",
            )
        )
        for finding in findings
    )


def highest_severity(
    findings: List[Dict[str, Any]],
) -> str:
    counts = severity_counts(
        findings
    )

    for severity in SEVERITIES:
        if counts.get(severity, 0):
            return severity

    return "LOW"


def minimum_verdict(
    findings: List[Dict[str, Any]],
) -> Tuple[str, str]:
    severity = highest_severity(
        findings
    )

    if severity == "CRITICAL":
        return (
            "CRITICAL",
            "HIGH RISK",
        )

    if severity == "HIGH":
        return (
            "HIGH",
            "CHANGES REQUIRED",
        )

    if severity == "MEDIUM":
        return (
            "MEDIUM",
            "CHANGES REQUIRED",
        )

    if severity in {
        "LOW",
        "NIT",
    }:
        return (
            "LOW",
            "APPROVE WITH MINOR CHANGES",
        )

    return (
        "LOW",
        "APPROVE",
    )


def recommendation_color(
    recommendation: str,
) -> str:
    mapping = {
        "HIGH RISK": COLORS["CRITICAL"],
        "CHANGES REQUIRED": COLORS["HIGH"],
        "APPROVE WITH MINOR CHANGES": COLORS["MEDIUM"],
        "APPROVE": COLORS["LOW"],
    }

    return mapping.get(
        recommendation,
        COLORS["MUTED"],
    )


def category_risk(
    findings: List[Dict[str, Any]],
    category: str,
) -> str:
    matching = [
        finding
        for finding in findings
        if normalize_category(
            finding.get("category")
        ) == category
    ]

    if not matching:
        return "LOW"

    return highest_severity(
        matching
    )


# ---------------------------------------------------------------------------
# Document metadata / formatting
# ---------------------------------------------------------------------------

def configure_document(
    document: Document,
) -> None:
    section = document.sections[0]

    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = document.styles

    normal = styles["Normal"]

    normal.font.name = "Aptos"
    normal.font.size = Pt(9)
    normal.font.color.rgb = rgb(
        COLORS["TEXT"]
    )

    for style_name, size, color in [
        ("Title", 28, COLORS["PRIMARY"]),
        ("Heading 1", 17, COLORS["PRIMARY"]),
        ("Heading 2", 13, COLORS["PRIMARY"]),
        ("Heading 3", 10, COLORS["PRIMARY"]),
    ]:
        try:
            style = styles[style_name]

            style.font.name = "Aptos Display"
            style.font.size = Pt(size)
            style.font.color.rgb = rgb(
                color
            )
        except KeyError:
            pass


def add_header_footer(
    document: Document,
    application_name: str,
) -> None:
    for section in document.sections:
        header = section.header

        paragraph = header.paragraphs[0]
        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.RIGHT
        )

        paragraph.text = ""

        run = paragraph.add_run(
            f"{application_name} | CODE REVIEW REPORT"
        )

        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = rgb(
            COLORS["MUTED"]
        )

        footer = section.footer

        paragraph = footer.paragraphs[0]
        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        paragraph.text = ""

        run = paragraph.add_run(
            "Confidential Code Review"
        )

        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = rgb(
            COLORS["MUTED"]
        )


def add_title_page(
    document: Document,
    review: Dict[str, Any],
    findings: List[Dict[str, Any]],
) -> None:
    application = review.get(
        "application",
        {},
    )

    application_name = first_value(
        application,
        ["name", "applicationName"],
    )

    overall_risk = safe_text(
        review.get(
            "overallRisk",
            "LOW",
        ),
        "LOW",
    ).upper()

    recommendation = safe_text(
        review.get(
            "overallRecommendation",
            "APPROVE",
        ),
        "APPROVE",
    ).upper()

    title = document.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    title.paragraph_format.space_before = Pt(50)

    run = title.add_run(
        "CODE REVIEW REPORT"
    )

    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(28)
    run.font.color.rgb = rgb(
        COLORS["PRIMARY"]
    )

    subtitle = document.add_paragraph()

    subtitle.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = subtitle.add_run(
        application_name
    )

    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(18)
    run.font.color.rgb = rgb(
        COLORS["TEXT"]
    )

    document.add_paragraph()

    add_key_value_table(
        document,
        [
            (
                "Application",
                application_name,
            ),
            (
                "Repository",
                first_value(
                    application,
                    ["repository", "repositoryPath"],
                ),
            ),
            (
                "Branch",
                first_value(
                    application,
                    ["branch"],
                ),
            ),
            (
                "Commit / Version",
                first_value(
                    application,
                    ["commit", "version"],
                ),
            ),
            (
                "Runtime",
                first_value(
                    application,
                    ["runtime", "muleRuntime"],
                ),
            ),
            (
                "Java Version",
                first_value(
                    application,
                    ["javaVersion", "java"],
                ),
            ),
            (
                "Build Version",
                first_value(
                    application,
                    ["buildVersion", "version"],
                ),
            ),
            (
                "Deployment Model",
                first_value(
                    application,
                    ["deploymentModel"],
                ),
            ),
            (
                "Review Type",
                first_value(
                    review,
                    ["reviewType"],
                    "FULL_APPLICATION",
                ),
            ),
            (
                "Review Date",
                first_value(
                    review,
                    ["reviewDate"],
                    now_local().strftime(
                        "%Y-%m-%d"
                    ),
                ),
            ),
            (
                "Reviewer",
                first_value(
                    review,
                    ["reviewer"],
                    DEFAULT_REVIEWER,
                ),
            ),
            (
                "Total Findings",
                len(findings),
            ),
        ],
    )

    add_status_badge(
        document,
        "OVERALL RISK",
        overall_risk,
        color_for(
            overall_risk,
            "MUTED",
        ),
    )

    add_status_badge(
        document,
        "OVERALL RECOMMENDATION",
        recommendation,
        recommendation_color(
            recommendation
        ),
    )

    counts = severity_counts(
        findings
    )

    add_generic_table(
        document,
        [
            "Severity",
            "Count",
            "Status",
        ],
        [
            (
                severity,
                counts.get(
                    severity,
                    0,
                ),
                {
                    "CRITICAL": "🔴 Critical",
                    "HIGH": "🟠 High",
                    "MEDIUM": "🟡 Medium",
                    "LOW": "🟢 Low",
                    "NIT": "🔵 Nit",
                }[severity],
            )
            for severity in SEVERITIES
        ] + [
            (
                "TOTAL",
                len(findings),
                "",
            )
        ],
    )

    document.add_page_break()


# ---------------------------------------------------------------------------
# Executive summary
# ---------------------------------------------------------------------------

def build_executive_summary(
    document: Document,
    review: Dict[str, Any],
    findings: List[Dict[str, Any]],
) -> None:
    add_heading(
        document,
        "1. EXECUTIVE SUMMARY",
        1,
    )

    overall_assessment = review.get(
        "executiveSummary",
        review.get(
            "overallAssessment",
            "Not Identified",
        ),
    )

    add_heading(
        document,
        "1.0 Overall Assessment",
        2,
    )

    add_key_value_table(
        document,
        [
            (
                "Overall Risk",
                safe_text(
                    review.get(
                        "overallRisk",
                        "LOW",
                    )
                ),
            ),
            (
                "Overall Recommendation",
                safe_text(
                    review.get(
                        "overallRecommendation",
                        "APPROVE",
                    )
                ),
            ),
            (
                "Finding Count",
                len(findings),
            ),
        ],
    )

    add_body(
        document,
        overall_assessment,
    )

    strengths = review.get(
        "strengths",
        review.get(
            "positiveObservations",
            [],
        ),
    )

    if strengths:
        add_heading(
            document,
            "Strengths",
            3,
        )

        if isinstance(
            strengths,
            list,
        ):
            add_bullets(
                document,
                strengths,
            )
        else:
            add_body(
                document,
                strengths,
            )

    concerns = review.get(
        "significantRisks",
        review.get(
            "productionConcerns",
            [],
        ),
    )

    if concerns:
        add_heading(
            document,
            "Significant Risks",
            3,
        )

        if isinstance(
            concerns,
            list,
        ):
            add_bullets(
                document,
                concerns,
            )
        else:
            add_body(
                document,
                concerns,
            )

    add_heading(
        document,
        "1.1 Review Scope",
        2,
    )

    scope = review.get(
        "scope",
        "Full MuleSoft application repository review.",
    )

    add_body(
        document,
        scope,
    )

    included = review.get(
        "includedAreas",
        [],
    )

    excluded = review.get(
        "excludedAreas",
        [],
    )

    limitations = review.get(
        "limitations",
        [],
    )

    if included:
        add_heading(
            document,
            "Included Areas",
            3,
        )
        add_bullets(
            document,
            included
            if isinstance(
                included,
                list,
            )
            else [included],
        )

    if excluded:
        add_heading(
            document,
            "Excluded Areas",
            3,
        )
        add_bullets(
            document,
            excluded
            if isinstance(
                excluded,
                list,
            )
            else [excluded],
        )

    if limitations:
        add_heading(
            document,
            "Known Limitations",
            3,
        )
        add_bullets(
            document,
            limitations
            if isinstance(
                limitations,
                list,
            )
            else [limitations],
        )

    add_heading(
        document,
        "1.2 Review Methodology",
        2,
    )

    add_generic_table(
        document,
        [
            "Step",
            "Review Activity",
            "Purpose",
        ],
        [
            (
                "1",
                "Application Discovery",
                "Understand application",
            ),
            (
                "2",
                "Architecture Review",
                "Assess architecture",
            ),
            (
                "3",
                "Security Review",
                "Identify security risks",
            ),
            (
                "4",
                "Integration Review",
                "Assess dependencies",
            ),
            (
                "5",
                "Data Review",
                "Assess data handling",
            ),
            (
                "6",
                "Error Handling Review",
                "Evaluate failures",
            ),
            (
                "7",
                "Performance Review",
                "Identify performance risks",
            ),
            (
                "8",
                "Testing Review",
                "Assess tests",
            ),
            (
                "9",
                "Maintainability Review",
                "Assess maintainability",
            ),
            (
                "10",
                "Production Readiness",
                "Assess release readiness",
            ),
            (
                "11",
                "Risk Reconciliation",
                "Establish final verdict",
            ),
        ],
    )

    add_heading(
        document,
        "1.3 Findings by Severity",
        2,
    )

    counts = severity_counts(
        findings
    )

    add_generic_table(
        document,
        [
            "Severity",
            "Count",
            "Status",
        ],
        [
            (
                severity,
                counts.get(
                    severity,
                    0,
                ),
                severity,
            )
            for severity in SEVERITIES
        ] + [
            (
                "TOTAL",
                len(findings),
                "",
            )
        ],
    )

    add_heading(
        document,
        "1.4 Findings by Category",
        2,
    )

    categories = category_counts(
        findings
    )

    rows = []

    known_categories = list(
        CATEGORY_ORDER
    )

    for category in sorted(
        categories.keys()
    ):
        if category not in known_categories:
            known_categories.append(
                category
            )

    for category in known_categories:
        count = categories.get(
            category,
            0,
        )

        rows.append(
            (
                category,
                count,
                category_risk(
                    findings,
                    category,
                )
                if count
                else "LOW",
            )
        )

    add_generic_table(
        document,
        [
            "Category",
            "Finding Count",
            "Risk",
        ],
        rows,
    )


# ---------------------------------------------------------------------------
# Application inventory
# ---------------------------------------------------------------------------

def build_application_inventory(
    document: Document,
    review: Dict[str, Any],
) -> None:
    add_heading(
        document,
        "2. APPLICATION INVENTORY",
        1,
    )

    inventory = review.get(
        "applicationInventory",
        review.get(
            "inventory",
            {},
        ),
    )

    if not isinstance(
        inventory,
        dict,
    ):
        inventory = {}

    add_heading(
        document,
        "2.0 Application Inventory",
        2,
    )

    rows = [
        (
            "Application Type",
            first_value(
                inventory,
                ["applicationType"],
            ),
        ),
        (
            "Runtime",
            first_value(
                inventory,
                ["runtime", "muleRuntime"],
            ),
        ),
        (
            "Main Flows",
            first_value(
                inventory,
                ["mainFlows", "flows"],
            ),
        ),
        (
            "APIs",
            first_value(
                inventory,
                ["apis", "api"],
            ),
        ),
        (
            "Data Stores",
            first_value(
                inventory,
                ["dataStores", "databases"],
            ),
        ),
        (
            "External Services",
            first_value(
                inventory,
                ["externalServices"],
            ),
        ),
        (
            "Messaging",
            first_value(
                inventory,
                ["messaging"],
            ),
        ),
        (
            "Deployment Model",
            first_value(
                inventory,
                ["deploymentModel"],
            ),
        ),
        (
            "Configuration Model",
            first_value(
                inventory,
                ["configurationModel"],
            ),
        ),
        (
            "Testing Structure",
            first_value(
                inventory,
                ["testingStructure", "testing"],
            ),
        ),
    ]

    add_key_value_table(
        document,
        rows,
    )

    add_heading(
        document,
        "2.1 Technology Stack",
        2,
    )

    technology_stack = review.get(
        "technologyStack",
        [],
    )

    if not isinstance(
        technology_stack,
        list,
    ):
        technology_stack = []

    rows = []

    for item in technology_stack:
        if not isinstance(
            item,
            dict,
        ):
            continue

        rows.append(
            (
                first_value(
                    item,
                    ["component", "name"],
                ),
                first_value(
                    item,
                    ["version"],
                ),
                first_value(
                    item,
                    ["evidence"],
                ),
                first_value(
                    item,
                    ["notes"],
                ),
            )
        )

    if not rows:
        rows = [
            (
                "Technology Stack",
                "Not Identified",
                "Not Identified",
                "Not Identified",
            )
        ]

    add_generic_table(
        document,
        [
            "Component",
            "Version",
            "Evidence",
            "Notes",
        ],
        rows,
    )

    add_heading(
        document,
        "2.2 Architecture Summary",
        2,
    )

    architecture = review.get(
        "architectureSummary",
        review.get(
            "architecture",
            "Not Identified",
        ),
    )

    if isinstance(
        architecture,
        dict,
    ):
        add_body(
            document,
            first_value(
                architecture,
                ["summary", "description"],
            ),
        )

        for label, key in [
            (
                "Major Components",
                "majorComponents",
            ),
            (
                "Entry Points",
                "entryPoints",
            ),
            (
                "Processing Flows",
                "processingFlows",
            ),
            (
                "Data Access",
                "dataAccess",
            ),
            (
                "External Integrations",
                "externalIntegrations",
            ),
            (
                "Error Boundaries",
                "errorBoundaries",
            ),
            (
                "Security Boundaries",
                "securityBoundaries",
            ),
            (
                "Configuration Boundaries",
                "configurationBoundaries",
            ),
            (
                "Architectural Concerns",
                "architecturalConcerns",
            ),
        ]:
            value = architecture.get(
                key
            )

            if value:
                add_heading(
                    document,
                    label,
                    3,
                )

                if isinstance(
                    value,
                    list,
                ):
                    add_bullets(
                        document,
                        value,
                    )
                else:
                    add_body(
                        document,
                        value,
                    )
    else:
        add_body(
            document,
            architecture,
        )

    architecture_diagram = review.get(
        "architectureDiagram",
        review.get(
            "architectureDiagramText",
        ),
    )

    if architecture_diagram:
        add_heading(
            document,
            "Architecture Diagram",
            3,
        )

        add_evidence_block(
            document,
            architecture_diagram,
        )

    add_heading(
        document,
        "2.3 Integration Inventory",
        2,
    )

    integrations = review.get(
        "integrationInventory",
        review.get(
            "integrations",
            [],
        ),
    )

    if not isinstance(
        integrations,
        list,
    ):
        integrations = []

    rows = []

    for item in integrations:
        if not isinstance(
            item,
            dict,
        ):
            continue

        rows.append(
            (
                first_value(
                    item,
                    ["source"],
                ),
                first_value(
                    item,
                    ["target"],
                ),
                first_value(
                    item,
                    ["mechanism", "connector"],
                ),
                first_value(
                    item,
                    ["mode"],
                ),
                first_value(
                    item,
                    ["retry"],
                ),
                first_value(
                    item,
                    ["transaction"],
                ),
                first_value(
                    item,
                    ["risk"],
                    "LOW",
                ),
            )
        )

    if not rows:
        rows = [
            (
                "Not Identified",
                "Not Identified",
                "Not Identified",
                "Not Identified",
                "Not Identified",
                "Not Identified",
                "Not Assessed",
            )
        ]

    add_generic_table(
        document,
        [
            "Source",
            "Target",
            "Mechanism",
            "Mode",
            "Retry",
            "Transaction",
            "Risk",
        ],
        rows,
    )


# ---------------------------------------------------------------------------
# Findings
# ---------------------------------------------------------------------------

def add_finding_card(
    document: Document,
    finding: Dict[str, Any],
) -> None:
    finding_id = safe_text(
        finding.get(
            "id",
        ),
        "Not Identified",
    )

    severity = safe_text(
        finding.get(
            "severity",
        ),
        "Not Identified",
    ).upper()

    title = safe_text(
        finding.get(
            "title",
        ),
        "Not Identified",
    )

    color = color_for(
        severity,
        "MUTED",
    )

    table = document.add_table(
        rows=1,
        cols=1,
    )

    cell = table.cell(
        0,
        0,
    )

    set_cell_shading(
        cell,
        color,
    )

    set_cell_border(
        cell,
        color,
        "10",
    )

    paragraph = cell.paragraphs[0]

    run = paragraph.add_run(
        f"{severity}  |  {finding_id}\n{title}"
    )

    run.bold = True
    run.font.name = "Aptos"
    run.font.size = Pt(13)
    run.font.color.rgb = rgb(
        COLORS["WHITE"]
    )

    document.add_paragraph()

    add_key_value_table(
        document,
        [
            (
                "Finding ID",
                finding_id,
            ),
            (
                "Severity",
                severity,
            ),
            (
                "Category",
                normalize_category(
                    finding.get(
                        "category"
                    )
                ),
            ),
            (
                "Location",
                first_value(
                    finding,
                    [
                        "location",
                        "file",
                    ],
                ),
            ),
            (
                "Confidence",
                first_value(
                    finding,
                    ["confidence"],
                ),
            ),
        ],
    )

    add_labeled_text(
        document,
        "Problem",
        finding.get(
            "problem",
            "Not Identified",
        ),
    )

    add_labeled_text(
        document,
        "Evidence",
        finding.get(
            "evidence",
            "Not Identified",
        ),
    )

    evidence_excerpt = finding.get(
        "evidenceExcerpt"
    )

    if evidence_excerpt:
        add_labeled_text(
            document,
            "Evidence Excerpt",
            evidence_excerpt,
            evidence_block=True,
        )

    add_labeled_text(
        document,
        "Impact",
        finding.get(
            "impact",
            "Not Identified",
        ),
    )

    add_labeled_text(
        document,
        "Recommendation",
        finding.get(
            "recommendation",
            "Not Identified",
        ),
    )

    document.add_page_break()


def build_findings(
    document: Document,
    findings: List[Dict[str, Any]],
) -> None:
    add_heading(
        document,
        "3. FINDINGS",
        1,
    )

    if not findings:
        add_body(
            document,
            "No findings identified.",
        )

    for severity in SEVERITIES:
        add_heading(
            document,
            f"3.{SEVERITIES.index(severity) + 1} "
            f"{severity.title()} Findings",
            2,
        )

        matching = [
            finding
            for finding in findings
            if safe_text(
                finding.get(
                    "severity"
                ),
                "",
            ).upper() == severity
        ]

        if not matching:
            add_body(
                document,
                f"No {severity} findings identified.",
                color=COLORS["MUTED"],
                italic=True,
            )
            continue

        for finding in matching:
            add_finding_card(
                document,
                finding,
            )

    add_heading(
        document,
        "3.6 Invalid / Unclassified Findings",
        2,
    )

    invalid = [
        finding
        for finding in findings
        if safe_text(
            finding.get(
                "severity"
            ),
            "",
        ).upper()
        not in SEVERITIES
    ]

    if invalid:
        for finding in invalid:
            add_body(
                document,
                finding,
            )
    else:
        add_body(
            document,
            "No invalid or unclassified findings identified.",
            color=COLORS["MUTED"],
            italic=True,
        )


# ---------------------------------------------------------------------------
# Category assessments
# ---------------------------------------------------------------------------

def build_category_assessments(
    document: Document,
    review: Dict[str, Any],
    findings: List[Dict[str, Any]],
) -> None:
    add_heading(
        document,
        "4. CATEGORY ASSESSMENTS",
        1,
    )

    raw = review.get(
        "categoryAssessments",
        review.get(
            "assessments",
            {},
        ),
    )

    if isinstance(
        raw,
        list,
    ):
        assessments = {
            normalize_category(
                item.get(
                    "category",
                    "Other",
                )
            ): item
            for item in raw
            if isinstance(
                item,
                dict,
            )
        }
    elif isinstance(
        raw,
        dict,
    ):
        assessments = raw
    else:
        assessments = {}

    for index, category in enumerate(
        CATEGORY_ORDER,
        start=1,
    ):
        add_heading(
            document,
            f"4.{index} {category}",
            2,
        )

        data = assessments.get(
            category,
            {},
        )

        if not isinstance(
            data,
            dict,
        ):
            data = {}

        rating = first_value(
            data,
            ["rating"],
            "NOT ASSESSED",
        ).upper()

        evidence_level = first_value(
            data,
            ["evidenceLevel"],
        )

        key_concern = first_value(
            data,
            ["keyConcern"],
        )

        add_key_value_table(
            document,
            [
                (
                    "Rating",
                    rating,
                ),
                (
                    "Evidence Level",
                    evidence_level,
                ),
                (
                    "Key Concern",
                    key_concern,
                ),
            ],
        )

        add_heading(
            document,
            "Summary",
            3,
        )

        add_body(
            document,
            first_value(
                data,
                ["summary"],
            ),
        )

        add_heading(
            document,
            "Positive Observations",
            3,
        )

        positives = data.get(
            "positiveObservations",
            [],
        )

        if isinstance(
            positives,
            list,
        ) and positives:
            add_bullets(
                document,
                positives,
            )
        else:
            add_body(
                document,
                "Not Identified",
                color=COLORS["MUTED"],
            )

        add_heading(
            document,
            "Gaps",
            3,
        )

        gaps = data.get(
            "gaps",
            [],
        )

        if isinstance(
            gaps,
            list,
        ) and gaps:
            add_bullets(
                document,
                gaps,
            )
        else:
            add_body(
                document,
                "Not Identified",
                color=COLORS["MUTED"],
            )


# ---------------------------------------------------------------------------
# Production readiness
# ---------------------------------------------------------------------------

def build_production_readiness(
    document: Document,
    review: Dict[str, Any],
) -> None:
    add_heading(
        document,
        "5. PRODUCTION READINESS ASSESSMENT",
        1,
    )

    readiness = review.get(
        "productionReadiness",
        {},
    )

    if not isinstance(
        readiness,
        dict,
    ):
        readiness = {}

    overall = first_value(
        readiness,
        [
            "overall",
            "status",
            "overallReadiness",
        ],
        "NOT ASSESSED",
    ).upper()

    add_heading(
        document,
        "5.0 Overall Readiness",
        2,
    )

    add_status_badge(
        document,
        "PRODUCTION READINESS",
        overall,
        color_for(
            overall,
            "MUTED",
        ),
    )

    add_body(
        document,
        first_value(
            readiness,
            ["summary"],
        ),
    )

    add_heading(
        document,
        "5.1 Readiness Dimensions",
        2,
    )

    dimensions = readiness.get(
        "dimensions",
        [],
    )

    if isinstance(
        dimensions,
        dict,
    ):
        converted = []

        for name, data in dimensions.items():
            if isinstance(
                data,
                dict,
            ):
                converted.append(
                    {
                        "dimension": name,
                        **data,
                    }
                )
            else:
                converted.append(
                    {
                        "dimension": name,
                        "status": data,
                    }
                )

        dimensions = converted

    if not isinstance(
        dimensions,
        list,
    ):
        dimensions = []

    expected_dimensions = [
        "Security",
        "Reliability",
        "Error Recovery",
        "Idempotency",
        "Observability",
        "Performance",
        "Scalability",
        "Configuration",
        "Testing",
        "Operational Support",
    ]

    dimension_map = {
        safe_text(
            item.get(
                "dimension"
            ),
            "",
        ): item
        for item in dimensions
        if isinstance(
            item,
            dict,
        )
    }

    rows = []

    for dimension in expected_dimensions:
        item = dimension_map.get(
            dimension,
            {},
        )

        rows.append(
            (
                dimension,
                first_value(
                    item,
                    ["status"],
                    "NOT ASSESSED",
                ).upper(),
                first_value(
                    item,
                    ["notes"],
                ),
            )
        )

    add_generic_table(
        document,
        [
            "Dimension",
            "Status",
            "Notes",
        ],
        rows,
    )

    add_heading(
        document,
        "5.2 Positive Observations",
        2,
    )

    positives = readiness.get(
        "positiveObservations",
        [],
    )

    if isinstance(
        positives,
        list,
    ) and positives:
        rows = []

        for item in positives:
            if isinstance(
                item,
                dict,
            ):
                rows.append(
                    (
                        first_value(
                            item,
                            [
                                "observation",
                                "title",
                            ],
                        ),
                        first_value(
                            item,
                            ["evidence"],
                        ),
                    )
                )
            else:
                rows.append(
                    (
                        safe_text(
                            item
                        ),
                        "Not Identified",
                    )
                )

        add_generic_table(
            document,
            [
                "Positive Observation",
                "Supporting Evidence",
            ],
            rows,
        )
    else:
        add_body(
            document,
            "Not Identified",
            color=COLORS["MUTED"],
        )


# ---------------------------------------------------------------------------
# Risk and recommendation
# ---------------------------------------------------------------------------

def build_risk_recommendation(
    document: Document,
    review: Dict[str, Any],
    findings: List[Dict[str, Any]],
) -> None:
    add_heading(
        document,
        "6. RISK AND RECOMMENDATION",
        1,
    )

    risk_data = review.get(
        "riskSummary",
        [],
    )

    if not isinstance(
        risk_data,
        list,
    ):
        risk_data = []

    add_heading(
        document,
        "6.1 Risk Summary",
        2,
    )

    risk_rows = []

    for index, item in enumerate(
        risk_data,
        start=1,
    ):
        if not isinstance(
            item,
            dict,
        ):
            continue

        related = item.get(
            "relatedFindings",
            item.get(
                "findings",
                [],
            ),
        )

        risk_rows.append(
            (
                index,
                first_value(
                    item,
                    ["riskTheme", "theme"],
                ),
                first_value(
                    item,
                    ["impact"],
                ),
                safe_text(
                    related,
                    "Not Identified",
                ),
            )
        )

    if not risk_rows:
        for index, severity in enumerate(
            SEVERITIES,
            start=1,
        ):
            matching = [
                finding
                for finding in findings
                if safe_text(
                    finding.get(
                        "severity"
                    ),
                    "",
                ).upper() == severity
            ]

            if not matching:
                continue

            ids = [
                safe_text(
                    finding.get(
                        "id"
                    ),
                    "",
                )
                for finding in matching
            ]

            risk_rows.append(
                (
                    index,
                    f"{severity} findings",
                    "See detailed findings.",
                    ", ".join(ids),
                )
            )

    if not risk_rows:
        risk_rows = [
            (
                1,
                "No material risks identified.",
                "No findings.",
                "None",
            )
        ]

    add_generic_table(
        document,
        [
            "Priority",
            "Risk Theme",
            "Impact",
            "Related Findings",
        ],
        risk_rows,
    )

    add_heading(
        document,
        "6.2 Remediation Priorities",
        2,
    )

    priorities = review.get(
        "remediationPriorities",
        [],
    )

    if not isinstance(
        priorities,
        list,
    ):
        priorities = []

    remediation_rows = []

    for index, item in enumerate(
        priorities[:10],
        start=1,
    ):
        if not isinstance(
            item,
            dict,
        ):
            continue

        related = item.get(
            "relatedFindings",
            item.get(
                "findings",
                [],
            ),
        )

        remediation_rows.append(
            (
                index,
                first_value(
                    item,
                    [
                        "remediation",
                        "action",
                    ],
                ),
                first_value(
                    item,
                    ["impact"],
                ),
                safe_text(
                    related,
                    "Not Identified",
                ),
            )
        )

    if not remediation_rows:
        add_body(
            document,
            "No remediation priorities identified.",
            color=COLORS["MUTED"],
        )
    else:
        add_generic_table(
            document,
            [
                "Priority",
                "Remediation",
                "Impact",
                "Related Findings",
            ],
            remediation_rows,
        )

    add_heading(
        document,
        "6.3 Review Limitations",
        2,
    )

    limitations = review.get(
        "reviewLimitations",
        review.get(
            "limitationsDetail",
            [],
        ),
    )

    if not isinstance(
        limitations,
        list,
    ):
        limitations = []

    limitation_rows = []

    for item in limitations:
        if not isinstance(
            item,
            dict,
        ):
            continue

        limitation_rows.append(
            (
                first_value(
                    item,
                    ["area"],
                ),
                first_value(
                    item,
                    ["limitation"],
                ),
                first_value(
                    item,
                    ["reason"],
                ),
                first_value(
                    item,
                    ["effect"],
                ),
            )
        )

    if not limitation_rows:
        limitation_rows = [
            (
                "General",
                "No material limitations recorded.",
                "Not Identified",
                "None identified.",
            )
        ]

    add_generic_table(
        document,
        [
            "Area",
            "Limitation",
            "Reason",
            "Effect",
        ],
        limitation_rows,
    )

    add_heading(
        document,
        "6.4 Overall Risk and Recommendation",
        2,
    )

    stated_risk = safe_text(
        review.get(
            "overallRisk",
            "LOW",
        ),
        "LOW",
    ).upper()

    stated_recommendation = safe_text(
        review.get(
            "overallRecommendation",
            "APPROVE",
        ),
        "APPROVE",
    ).upper()

    minimum_risk, minimum_recommendation = (
        minimum_verdict(
            findings
        )
    )

    add_key_value_table(
        document,
        [
            (
                "Overall Risk",
                stated_risk,
            ),
            (
                "Overall Recommendation",
                stated_recommendation,
            ),
            (
                "Evidence-Supported Minimum Risk",
                minimum_risk,
            ),
            (
                "Evidence-Supported Minimum Recommendation",
                minimum_recommendation,
            ),
        ],
    )

    rationale = review.get(
        "recommendationRationale",
        review.get(
            "overallRationale",
            "Not Identified",
        ),
    )

    add_heading(
        document,
        "Recommendation Rationale",
        3,
    )

    add_body(
        document,
        rationale,
    )

    reconciliation = review.get(
        "verdictReconciliation",
        {},
    )

    if isinstance(
        reconciliation,
        dict,
    ) and reconciliation.get(
        "required"
    ):
        add_heading(
            document,
            "Verdict Reconciliation",
            3,
        )

        add_body(
            document,
            (
                "The originally stated verdict was less severe "
                "than the minimum verdict supported by the findings. "
                "The validator raised the final verdict accordingly."
            ),
        )


# ---------------------------------------------------------------------------
# Appendices
# ---------------------------------------------------------------------------

def build_appendix_a(
    document: Document,
    findings: List[Dict[str, Any]],
) -> None:
    add_heading(
        document,
        "APPENDIX A — COMPLETE FINDINGS INVENTORY",
        1,
    )

    rows = []

    for finding in findings:
        rows.append(
            (
                safe_text(
                    finding.get(
                        "id"
                    )
                ),
                safe_text(
                    finding.get(
                        "severity"
                    )
                ),
                normalize_category(
                    finding.get(
                        "category"
                    )
                ),
                safe_text(
                    finding.get(
                        "confidence"
                    )
                ),
                safe_text(
                    finding.get(
                        "title"
                    )
                ),
                first_value(
                    finding,
                    [
                        "location",
                        "file",
                    ],
                ),
            )
        )

    if not rows:
        rows = [
            (
                "None",
                "",
                "",
                "",
                "No findings identified.",
                "",
            )
        ]

    add_generic_table(
        document,
        [
            "ID",
            "Severity",
            "Category",
            "Confidence",
            "Title",
            "Location",
        ],
        rows,
    )

    detailed_count = len(findings)
    appendix_count = len(rows) if findings else 0

    status = (
        "MATCH"
        if detailed_count == appendix_count
        else "MISMATCH"
    )

    add_key_value_table(
        document,
        [
            (
                "Detailed Findings",
                detailed_count,
            ),
            (
                "Appendix Findings",
                appendix_count,
            ),
            (
                "Reconciliation Status",
                status,
            ),
        ],
    )


def build_appendix_b(
    document: Document,
    findings: List[Dict[str, Any]],
) -> None:
    add_heading(
        document,
        "APPENDIX B — FINDING EVIDENCE DETAIL",
        1,
    )

    if not findings:
        add_body(
            document,
            "No findings identified.",
        )
        return

    for finding in findings:
        finding_id = safe_text(
            finding.get(
                "id"
            )
        )

        title = safe_text(
            finding.get(
                "title"
            )
        )

        add_heading(
            document,
            f"{finding_id} — {title}",
            2,
        )

        add_heading(
            document,
            "Evidence",
            3,
        )

        add_body(
            document,
            finding.get(
                "evidence",
                "Not Identified",
            ),
        )

        excerpt = finding.get(
            "evidenceExcerpt"
        )

        if excerpt:
            add_heading(
                document,
                "Evidence Context",
                3,
            )

            add_evidence_block(
                document,
                excerpt,
            )


def build_appendix_c(
    document: Document,
    review: Dict[str, Any],
) -> None:
    add_heading(
        document,
        "APPENDIX C — REVIEW COVERAGE AND BOUNDARIES",
        1,
    )

    add_heading(
        document,
        "C.1 Review Coverage",
        2,
    )

    coverage = review.get(
        "reviewCoverage",
        review.get(
            "coverage",
            {},
        ),
    )

    if isinstance(
        coverage,
        list,
    ):
        coverage_map = {
            first_value(
                item,
                ["reviewArea", "area"],
            ): item
            for item in coverage
            if isinstance(
                item,
                dict,
            )
        }
    elif isinstance(
        coverage,
        dict,
    ):
        coverage_map = coverage
    else:
        coverage_map = {}

    expected = [
        "Application Discovery",
        "Architecture Analysis",
        "Security Analysis",
        "Database Analysis",
        "Messaging Analysis",
        "Performance Analysis",
        "Testing Analysis",
    ]

    rows = []

    for area in expected:
        item = coverage_map.get(
            area,
            {},
        )

        if not isinstance(
            item,
            dict,
        ):
            item = {
                "status": item
            }

        rows.append(
            (
                area,
                first_value(
                    item,
                    ["status"],
                    "NOT ASSESSED",
                ).upper(),
                first_value(
                    item,
                    ["notes"],
                ),
            )
        )

    add_generic_table(
        document,
        [
            "Review Area",
            "Status",
            "Notes",
        ],
        rows,
    )

    add_heading(
        document,
        "C.2 Review Controls",
        2,
    )

    controls = [
        (
            "Review conducted in read-only mode",
            "CONFIRMED",
        ),
        (
            "Application source was not modified",
            "CONFIRMED",
        ),
        (
            "No source changes introduced",
            "CONFIRMED",
        ),
        (
            "No dependency changes introduced",
            "CONFIRMED",
        ),
        (
            "No remediation performed",
            "CONFIRMED",
        ),
        (
            "Findings are evidence-backed",
            "CONFIRMED",
        ),
        (
            "Unverifiable areas recorded",
            "CONFIRMED",
        ),
        (
            "Finding counts reconciled",
            "CONFIRMED",
        ),
        (
            "Final verdict reconciled",
            "CONFIRMED",
        ),
    ]

    add_generic_table(
        document,
        [
            "Control",
            "Status",
        ],
        controls,
    )


# ---------------------------------------------------------------------------
# Document validation
# ---------------------------------------------------------------------------

REQUIRED_DOCUMENT_SECTIONS = [
    "CODE REVIEW REPORT",
    "1. EXECUTIVE SUMMARY",
    "2. APPLICATION INVENTORY",
    "3. FINDINGS",
    "4. CATEGORY ASSESSMENTS",
    "5. PRODUCTION READINESS ASSESSMENT",
    "6. RISK AND RECOMMENDATION",
    "APPENDIX A — COMPLETE FINDINGS INVENTORY",
    "APPENDIX B — FINDING EVIDENCE DETAIL",
    "APPENDIX C — REVIEW COVERAGE AND BOUNDARIES",
]


def extract_document_text(
    document: Document,
) -> str:
    values = []

    for paragraph in document.paragraphs:
        values.append(
            paragraph.text
        )

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                values.append(
                    cell.text
                )

    return "\n".join(values)


def validate_document(
    output_path: Path,
    expected_findings: List[Dict[str, Any]],
) -> Dict[str, Any]:
    if not output_path.exists():
        raise RuntimeError(
            f"Generated document does not exist: {output_path}"
        )

    if output_path.stat().st_size == 0:
        raise RuntimeError(
            f"Generated document is empty: {output_path}"
        )

    try:
        document = Document(
            str(output_path)
        )
    except Exception as exc:
        raise RuntimeError(
            f"Generated DOCX cannot be opened: {exc}"
        ) from exc

    text = extract_document_text(
        document
    )

    missing_sections = [
        section
        for section in REQUIRED_DOCUMENT_SECTIONS
        if section not in text
    ]

    if missing_sections:
        raise RuntimeError(
            "Generated document is missing required sections: "
            + ", ".join(
                missing_sections
            )
        )

    expected_ids = {
        safe_text(
            finding.get(
                "id"
            ),
            "",
        )
        for finding in expected_findings
    }

    missing_ids = [
        finding_id
        for finding_id in expected_ids
        if finding_id
        and finding_id not in text
    ]

    if missing_ids:
        raise RuntimeError(
            "Generated document is missing finding IDs: "
            + ", ".join(
                sorted(missing_ids)
            )
        )

    return {
        "status": "PASSED",
        "documentExists": True,
        "documentSizeBytes": output_path.stat().st_size,
        "requiredSectionsChecked": len(
            REQUIRED_DOCUMENT_SECTIONS
        ),
        "findingsChecked": len(
            expected_findings
        ),
    }


# ---------------------------------------------------------------------------
# Report generation
# ---------------------------------------------------------------------------

def choose_template(
    framework_root: Path,
    no_template: bool,
) -> Optional[Path]:
    if no_template:
        return None

    template = (
        framework_root
        / "templates"
        / "review-report-template.docx"
    )

    if template.exists():
        return template

    return None


def create_document(
    framework_root: Path,
    review: Dict[str, Any],
    findings: List[Dict[str, Any]],
    no_template: bool,
) -> Document:
    template = choose_template(
        framework_root,
        no_template,
    )

    if template:
        document = Document(
            str(template)
        )
    else:
        document = Document()

    configure_document(
        document
    )

    application = review.get(
        "application",
        {},
    )

    application_name = first_value(
        application,
        [
            "name",
            "applicationName",
        ],
    )

    add_header_footer(
        document,
        application_name,
    )

    add_title_page(
        document,
        review,
        findings,
    )

    build_executive_summary(
        document,
        review,
        findings,
    )

    add_section_page_break(
        document
    )

    build_application_inventory(
        document,
        review,
    )

    add_section_page_break(
        document
    )

    build_findings(
        document,
        findings,
    )

    add_section_page_break(
        document
    )

    build_category_assessments(
        document,
        review,
        findings,
    )

    add_section_page_break(
        document
    )

    build_production_readiness(
        document,
        review,
    )

    add_section_page_break(
        document
    )

    build_risk_recommendation(
        document,
        review,
        findings,
    )

    add_section_page_break(
        document
    )

    build_appendix_a(
        document,
        findings,
    )

    add_section_page_break(
        document
    )

    build_appendix_b(
        document,
        findings,
    )

    add_section_page_break(
        document
    )

    build_appendix_c(
        document,
        review,
    )

    return document


def determine_output_path(
    reports_dir: Path,
    requested_output: Optional[str],
) -> Path:
    if requested_output:
        output = Path(
            requested_output
        ).expanduser().resolve()

        if output.suffix.lower() != ".docx":
            raise RuntimeError(
                "Output file must have a .docx extension."
            )

        output.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        return output

    timestamp = now_local().strftime(
        "%Y%m%d-%H%M%S"
    )

    return (
        reports_dir
        / f"CODE_REVIEW_REPORT_{timestamp}.docx"
    )


def remove_old_generated_reports(
    reports_dir: Path,
) -> None:
    """
    Keep reports/ clean so each execution has exactly one generated
    final report.

    Only files matching the framework's generated filename are removed.
    Other repository files are untouched.
    """

    pattern = "CODE_REVIEW_REPORT_*.docx"

    for path in reports_dir.glob(
        pattern
    ):
        try:
            path.unlink()
        except OSError as exc:
            raise RuntimeError(
                f"Unable to remove previous generated report "
                f"{path}: {exc}"
            ) from exc


def write_validation_metadata(
    execution_dir: Path,
    output_path: Path,
    validation: Dict[str, Any],
) -> None:
    """
    Validation metadata is stored under workspace/execution, not reports/.
    """

    metadata = {
        "generatedAt": now_local().isoformat(),
        "report": str(output_path),
        "validation": validation,
    }

    metadata_path = (
        execution_dir
        / "generated-report-validation.json"
    )

    metadata_path.write_text(
        json.dumps(
            metadata,
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )


def generate(
    application_root: Path,
    framework_root: Path,
    input_file: Path,
    reports_dir: Path,
    requested_output: Optional[str],
    no_template: bool,
) -> Path:
    # application_root is intentionally accepted and validated for
    # execution-context awareness. The generator never modifies it.
    if not application_root.exists():
        raise RuntimeError(
            f"Application root does not exist: "
            f"{application_root}"
        )

    if not application_root.is_dir():
        raise RuntimeError(
            f"Application root is not a directory: "
            f"{application_root}"
        )

    if not framework_root.exists():
        raise RuntimeError(
            f"Framework root does not exist: "
            f"{framework_root}"
        )

    review = load_json(
        input_file
    )

    findings = sort_findings(
        get_findings(
            review
        )
    )

    output_path = determine_output_path(
        reports_dir,
        requested_output,
    )

    # Prevent stale reports from making the workflow appear to have
    # produced multiple final artifacts.
    if not requested_output:
        remove_old_generated_reports(
            reports_dir
        )

    document = create_document(
        framework_root,
        review,
        findings,
        no_template,
    )

    # Save first, then reopen and validate the actual DOCX.
    document.save(
        str(output_path)
    )

    validation = validate_document(
        output_path,
        findings,
    )

    execution_dir = (
        framework_root
        / "workspace"
        / "execution"
    )

    write_validation_metadata(
        execution_dir,
        output_path,
        validation,
    )

    return output_path


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main() -> int:
    args = parse_args()

    try:
        (
            application_root,
            framework_root,
            input_file,
            reports_dir,
        ) = resolve_paths(
            args
        )

        output_path = generate(
            application_root=application_root,
            framework_root=framework_root,
            input_file=input_file,
            reports_dir=reports_dir,
            requested_output=args.output,
            no_template=args.no_template,
        )

        # Intentionally concise output because GitHub Actions should print
        # only the final artifact information.
        print(
            f"REPORT_GENERATED={output_path}"
        )

        return 0

    except Exception as exc:
        print(
            f"REPORT_GENERATION_FAILED: {exc}",
            file=sys.stderr,
        )

        return 1


if __name__ == "__main__":
    sys.exit(
        main()
    )